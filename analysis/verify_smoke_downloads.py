"""Compare downloaded production workbooks against local smoke calculations."""
import argparse
import json
from pathlib import Path
from openpyxl import load_workbook

def compare(local, downloaded):
    a=load_workbook(local,read_only=True,data_only=True)
    b=load_workbook(downloaded,read_only=True,data_only=True)
    assert a.sheetnames == b.sheetnames, (a.sheetnames,b.sheetnames)
    differences=[]
    for name in a.sheetnames:
        av=list(a[name].values);bv=list(b[name].values)
        assert len(av)==len(bv),(name,len(av),len(bv))
        for i,(ar,br) in enumerate(zip(av,bv),1):
            assert len(ar)==len(br),(name,i)
            for j,(ac,bc) in enumerate(zip(ar,br),1):
                if ac!=bc: differences.append({'sheet':name,'row':i,'column':j,'local':ac,'production':bc})
    a.close();b.close()
    assert not differences,differences
    return {'download':Path(downloaded).name,'identical_cells':True}

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--source-dir',required=True);p.add_argument('--recommendations',required=True);p.add_argument('--summary',required=True);a=p.parse_args()
    d=Path(a.source_dir)
    report=[compare(d/'make_excel_simple.xlsx',a.recommendations),compare(d/'make_strategy_summary_excel.xlsx',a.summary)]
    (d/'production_download_verification.json').write_text(json.dumps(report,indent=2,ensure_ascii=False),encoding='utf-8')
    print('Both production downloads match all local workbook cells.')
