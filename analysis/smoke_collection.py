"""Private synthetic upload fixture and real-ledger download smoke check."""
import argparse
import io
import json
import sys
from pathlib import Path
from openpyxl import Workbook, load_workbook
import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT), str(ROOT / 'tests')]
from test_v2_15_4 import load_app_namespace
import ledger_metadata
import bid_strategies

def run(directory):
    directory = Path(directory)
    wb = Workbook(); ws = wb.active
    ws.append(['Synthetic smoke fixture; no real submission'])
    ws.append(['번호','공고명','공고번호','기초금액','투찰마감','발주기관','업종','지역','낙찰하한율'])
    cases = [
        ('전기공사', '전기', '부산광역시', 150_000_000, '부산'),
        ('건설사업관리 용역', '건설사업관리', '부산광역시', 150_000_000, '전국'),
        ('설계용역', '설계', '부산광역시', 150_000_000, '전국'),
        ('전력시설물 감리용역', '감리', '한국전력공사 부산울산본부', 80_000_000, '부산'),
    ]
    for i, (name, industry, org, base, region) in enumerate(cases, 1):
        ws.append([i, name, f'SMOKE20261014-{i:03d}-00', base, '26.09.30 (10:00)', org, industry, region, 88.])
    sample = directory / 'sample_bid_folder.xlsx'; wb.save(sample)
    app = load_app_namespace()
    bids = app['parse_xls'](sample.read_bytes(), sample.name)
    assert len(bids) == 4
    history = pd.read_pickle(directory / 'operational.pkl')
    metadata = json.loads((directory / 'manifest.json').read_text(encoding='utf-8'))
    history.attrs[ledger_metadata.ATTR] = metadata
    prepared = app['prepare_history_frame'](history)
    results = app['predict_final_batch'](bids, prepared, strategy_mode='auto')
    for b, result in zip(bids, results):
        result.update(bid=b, scope_info=app['classify_bid_scope'](b))
    summary = []
    for func in ['make_excel_simple', 'make_strategy_summary_excel']:
        base = app[func](results, app['history_quality_summary'](history)) if func == 'make_excel_simple' else app[func](results)
        output = bid_strategies.add_review_sheets(ledger_metadata.add_excel_metadata(base, metadata), results, [])
        file = directory / (func + '.xlsx'); file.write_bytes(output.getvalue())
        book = load_workbook(file, read_only=True, data_only=True)
        assert '데이터출처' in book.sheetnames
        summary.append({'function':func, 'sheets':book.sheetnames, 'first_rows':list(book.worksheets[0].values)[:6]})
        book.close()
    result = {'bids':len(bids), 'training_rows':len(prepared), 'version':metadata['version'],
              'recommendation_counts':[len(r['recommendations']) for r in results],
              'rates':[[v['rate'] for v in r['recommendations']] for r in results], 'excels':summary}
    assert result['recommendation_counts'] == [1,3,3,2], result
    (directory / 'smoke.json').write_text(json.dumps(result,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
    print(json.dumps(result,ensure_ascii=True,default=str))

if __name__ == '__main__':
    p=argparse.ArgumentParser();p.add_argument('--source-dir',required=True)
    run(p.parse_args().source_dir)
