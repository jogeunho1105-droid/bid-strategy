"""Export the validated ledger with complete lineage and source-label reviews."""
import json
from pathlib import Path
import pandas as pd
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.cell import WriteOnlyCell
from openpyxl.utils import get_column_letter

import argparse
parser=argparse.ArgumentParser();parser.add_argument('--source-dir',required=True)
p=Path(parser.parse_args().source_dir);m=json.loads((p/'manifest.json').read_text(encoding='utf-8'))
d=pd.read_pickle(p/'operational.pkl');lineage=pd.read_csv(p/'source_excel_lineage.csv.gz',keep_default_na=False)
d=d.merge(lineage[['task','source_csv_row','source_excel','source_excel_row']],left_on=['source_task','source_csv_row'],right_on=['task','source_csv_row'],how='left',validate='one_to_one').drop(columns=['task'])
review=pd.read_csv(p/'classification_review.csv',dtype=str,keep_default_na=False)
needs_review=d.event_id.isin(set(review.event_id))
d.loc[needs_review,'quality_warnings']=d.loc[needs_review,'quality_warnings'].map(lambda v:(';'.join([v,'dual_profile_classification_review'])).strip(';'))
assert len(d)==m['operational_rows'] and d.source_excel.notna().all()
columns=m['original_columns']+[c for c in d if c not in m['original_columns']]
wb=Workbook(write_only=True);ws=wb.create_sheet('통합낙찰이력')
ws.freeze_panes='A2';ws.auto_filter.ref=f'A1:{get_column_letter(len(columns))}{len(d)+1}'
for i in range(1,len(columns)+1):ws.column_dimensions[get_column_letter(i)].width=22
ws.row_dimensions[1].height=32
for key,width in [('A',10),('B',65),('C',25),('D',42),('E',18),('F',18),('G',18),('O',16),('V',30),('AB',18)]:ws.column_dimensions[key].width=width
head=[]
for name in columns:
    cell=WriteOnlyCell(ws,value=name);cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='1A2744');cell.alignment=Alignment(wrap_text=True,vertical='center');head.append(cell)
ws.append(head)
for row in d[columns].itertuples(index=False,name=None):
    values=[]
    for i,v in enumerate(row):
        if pd.isna(v):values.append(None)
        elif isinstance(v,pd.Timestamp):
            cell=WriteOnlyCell(ws,value=v.to_pydatetime());cell.number_format='yyyy-mm-dd';values.append(cell)
        else:values.append(v)
    ws.append(values)
meta=wb.create_sheet('데이터출처');meta.column_dimensions['A'].width=30;meta.column_dimensions['B'].width=100
meta.append(['신규 기준 원장 출처정보 — 원본·기존 원장 보존']);meta.append([json.dumps(m,ensure_ascii=False)])
for row in [('버전',m['version']),('원천행',m['source_rows']),('운영 사건수',m['operational_rows']),('학습 후보',m['training_rows']),('학습 최종일',m['data_end_date']),('완료작업',f"{m['completed_tasks']}/{m['total_tasks']}"),('미해결',m['unresolved']),('지역정보',m['region_basis']),('숫자 품질상태',json.dumps(m['numeric_status'],ensure_ascii=False)),('품질경고',json.dumps(m['warnings'],ensure_ascii=False))]:
    a=WriteOnlyCell(meta,value=row[0]);a.font=Font(bold=True)
    b=WriteOnlyCell(meta,value=row[1]);b.alignment=Alignment(wrap_text=True);meta.append([a,b])
source=wb.create_sheet('원천행출처');source.freeze_panes='A2'
for i,col in enumerate(lineage.columns,1):source.column_dimensions[get_column_letter(i)].width=100 if col in ['source_csv','source_excel'] else 25
source.append(list(lineage.columns))
for row in lineage.itertuples(index=False,name=None):source.append(row)
classes=wb.create_sheet('분류확인');classes.freeze_panes='A2'
for i,col in enumerate(review.columns,1):classes.column_dimensions[get_column_letter(i)].width=65 if col=='공고명' else 100 if col.endswith('csv') else 35
classes.append(list(review.columns))
for row in review.itertuples(index=False,name=None):classes.append(row)
file=p/'신규_운영_낙찰데이터_20261001.xlsx';wb.save(file)
w=load_workbook(file,read_only=True,data_only=True)
assert w.sheetnames==['통합낙찰이력','데이터출처','원천행출처','분류확인']
total=sum(1 for _ in w['통합낙찰이력'].values)-1;assert total==112127
assert sum(1 for _ in w['원천행출처'].values)-1==112848
assert sum(1 for _ in w['분류확인'].values)-1==496
w.close();print('exported',file.name,total,file.stat().st_size,flush=True)

