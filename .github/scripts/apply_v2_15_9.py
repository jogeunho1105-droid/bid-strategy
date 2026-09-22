from pathlib import Path

ROOT = Path('.')
APP_PATH = ROOT / '입찰 앱.py'
README_PATH = ROOT / 'README.md'
TEST_V157_PATH = ROOT / 'tests' / 'test_app_v2_15_7.py'
TEST_V158_PATH = ROOT / 'tests' / 'test_app_v2_15_8.py'
TEST_V159_PATH = ROOT / 'tests' / 'test_app_v2_15_9.py'
VERIFY_PATH = ROOT / 'docs' / 'change_requests' / 'v2.15.9-implementation-verification.md'


def replace_once(text: str, old: str, new: str, label: str) -> str:
    if new in text:
        return text
    count = text.count(old)
    if count != 1:
        raise RuntimeError(f'{label}: expected one old value, found {count}')
    return text.replace(old, new, 1)


app = APP_PATH.read_text(encoding='utf-8')
app = replace_once(
    app,
    '# ║  투찰전략 분석 시스템 v2.15.8                                   ║',
    '# ║  투찰전략 분석 시스템 v2.15.9                                   ║',
    'header version',
)
app = replace_once(
    app,
    'MODEL_VERSION = "v2.15.8"',
    'MODEL_VERSION = "v2.15.9"',
    'model version',
)
app = replace_once(
    app,
    '<h2>📊 투찰전략 분석 시스템 v2.15.8</h2>',
    '<h2>📊 투찰전략 분석 시스템 v2.15.9</h2>',
    'screen version',
)

parse_anchor = 'def parse_xls(file_bytes, filename=""):\n'
parse_helper = '''def parse_lower_limit_rate(value):
    """입찰서류함 낙찰하한율을 퍼센트 숫자(예: 89.745)로 정규화한다."""
    if value is None:
        return None
    try:
        if isinstance(value, str):
            text=value.strip().replace(",", "")
            if not text:
                return None
            had_percent="%" in text
            text=text.replace("%", "").strip()
            number=float(text)
            if not had_percent and 0 < abs(number) <= 1:
                number*=100
        else:
            number=float(value)
            if not np.isfinite(number):
                return None
            if 0 < abs(number) <= 1:
                number*=100
        if not np.isfinite(number) or number <= 0 or number > 100:
            return None
        return round(float(number),6)
    except (TypeError,ValueError):
        return None


def parse_xls(file_bytes, filename=""):
'''
if 'def parse_lower_limit_rate(value):' not in app:
    if parse_anchor not in app:
        raise RuntimeError('parse_xls anchor not found')
    app = app.replace(parse_anchor, parse_helper, 1)

xlsx_old = '''                "companies": row.get("업체수") or row.get("참가업체수") or None,
            })'''
xlsx_new = '''                "companies": row.get("업체수") or row.get("참가업체수") or None,
                "lower_limit_rate": parse_lower_limit_rate(row.get("낙찰하한율")),
            })'''
if '"lower_limit_rate": parse_lower_limit_rate(row.get("낙찰하한율"))' not in app:
    app = replace_once(app, xlsx_old, xlsx_new, 'xlsx lower-limit parser')

xls_old = '''                "companies": row.get("업체수", row.get("참가업체수", None)),
            })'''
xls_new = '''                "companies": row.get("업체수", row.get("참가업체수", None)),
                "lower_limit_rate": parse_lower_limit_rate(row.get("낙찰하한율")),
            })'''
if app.count('"lower_limit_rate": parse_lower_limit_rate(row.get("낙찰하한율"))') < 2:
    app = replace_once(app, xls_old, xls_new, 'xls lower-limit parser')

main_ui_anchor = '# ════════════════════════════════════════════════════════════════\n#  메인 UI\n'
summary_function = '''def format_lower_limit_rate(value):
    rate=parse_lower_limit_rate(value)
    if rate is None:
        return "미확인"
    return f"{rate:.6f}".rstrip("0").rstrip(".")+"%"


def make_strategy_summary_excel(results):
    """공고명·낙찰하한율·업체1~3 최종 추천만 포함한 요약 Excel."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb=Workbook(); ws=wb.active; ws.title="투찰전략"
    ws.sheet_view.showGridLines=False
    navy="FF1a2744"; light="FFf8fafc"
    thin=Side(style="thin",color="FFd1d5db")
    border=Border(left=thin,right=thin,top=thin,bottom=thin)
    headers=["공고명","낙찰하한율","업체1추천","업체2추천","업체3추천"]
    widths=[60,15,15,15,15]

    ws.merge_cells("A1:E1")
    title=ws["A1"]
    title.value=f"투찰전략 분석 — {MODEL_VERSION} / {datetime.now().strftime('%Y.%m.%d')}"
    title.font=Font(name="맑은 고딕",bold=True,size=14,color="FFFFFFFF")
    title.fill=PatternFill("solid",start_color=navy)
    title.alignment=Alignment(horizontal="center",vertical="center")
    ws.row_dimensions[1].height=30

    for col,(header,width) in enumerate(zip(headers,widths),1):
        cell=ws.cell(2,col,header)
        cell.font=Font(name="맑은 고딕",bold=True,color="FFFFFFFF")
        cell.fill=PatternFill("solid",start_color=navy)
        cell.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
        cell.border=border
        ws.column_dimensions[get_column_letter(col)].width=width
    ws.row_dimensions[2].height=26

    for row_idx,row in enumerate(results,3):
        bid=row.get("bid") or {}
        recs=row.get("recommendations") or []
        scope_info=row.get("scope_info") or {}
        company_count=int(scope_info.get("company_count",len(recs) or 3))
        values=[str(bid.get("name") or ""),format_lower_limit_rate(bid.get("lower_limit_rate"))]
        for pos in range(3):
            if pos < len(recs):
                values.append(f"{float(recs[pos]['rate']):+.4f}%")
            elif pos >= company_count:
                values.append("참여대상 없음")
            else:
                values.append("없음")
        for col,value in enumerate(values,1):
            cell=ws.cell(row_idx,col,value)
            cell.font=Font(name="맑은 고딕",size=10)
            cell.fill=PatternFill("solid",start_color=light if row_idx%2==0 else "FFFFFFFF")
            cell.alignment=Alignment(
                horizontal="left" if col==1 else "center",
                vertical="center",
                wrap_text=(col==1),
            )
            cell.border=border
        ws.row_dimensions[row_idx].height=32

    last_row=max(2,len(results)+2)
    ws.auto_filter.ref=f"A2:E{last_row}"
    ws.freeze_panes="A3"
    buf=io.BytesIO(); wb.save(buf); buf.seek(0); return buf


'''
if 'def make_strategy_summary_excel(results):' not in app:
    if main_ui_anchor not in app:
        raise RuntimeError('main UI anchor not found')
    app = app.replace(main_ui_anchor, summary_function + main_ui_anchor, 1)

old_download = '''    st.download_button("📥 업체 추천표 다운로드",
        data=excel_buf,
        file_name=f"투찰전략_{today_str}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        type="primary",use_container_width=True)'''
new_download = '''    st.download_button("📥 업체 추천표 다운로드",
        data=excel_buf,
        file_name=f"투찰전략_{today_str}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        type="primary",use_container_width=True)
    summary_excel_buf=make_strategy_summary_excel(results)
    st.download_button("📥 투찰전략 요약 다운로드",
        data=summary_excel_buf,
        file_name=f"투찰전략_요약_{today_str}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True)'''
if '📥 투찰전략 요약 다운로드' not in app:
    app = replace_once(app, old_download, new_download, 'summary download button')

APP_PATH.write_text(app, encoding='utf-8')

readme = README_PATH.read_text(encoding='utf-8')
readme = replace_once(readme, '# 투찰전략 분석 시스템 v2.15.8', '# 투찰전략 분석 시스템 v2.15.9', 'README title')
readme = replace_once(
    readme,
    'v2.15.8은 사용자 승인에 따른 요약표 가독성 개선 배포 버전입니다.',
    'v2.15.9는 사용자 승인에 따라 낙찰하한율과 업체별 추천을 담은 투찰전략 요약 Excel 다운로드를 추가한 배포 버전입니다.',
    'README release sentence',
)
readme = replace_once(
    readme,
    'v2.15.8은 실행파일 하나만 교체해서는 동작하지 않습니다.',
    'v2.15.9는 실행파일 하나만 교체해서는 동작하지 않습니다.',
    'README deployment version',
)
readme = replace_once(
    readme,
    '5. 업체별 추천표를 Excel로 내려받고, 개찰 후 `사후낙찰검증` 시트에 당사 실제 투찰정보와 결과를 기록합니다.',
    '5. `투찰전략 요약 다운로드`로 공고명·낙찰하한율·업체별 추천만 내려받거나, `업체 추천표 다운로드`로 상세 Excel을 내려받아 개찰 후 `사후낙찰검증` 시트에 당사 실제 투찰정보와 결과를 기록합니다.',
    'README download usage',
)
section = '''## v2.15.9 투찰전략 요약 다운로드

입찰서류함의 `낙찰하한율`을 원본에서 읽어 다음 열만 포함한 별도 Excel을 제공합니다.

```text
공고명 | 낙찰하한율 | 업체1추천 | 업체2추천 | 업체3추천
```

`낙찰하한율`은 앱에서 재계산하거나 추정하지 않습니다. 숫자 또는 `%` 문자열을 퍼센트 숫자로 정규화하며, 누락·비정상 값은 `미확인`으로 표시합니다. 참여하지 않는 업체는 `참여대상 없음`, 추천 데이터가 부족하면 `없음`으로 표시합니다. 기존 상세 Excel과 화면 요약표, 추천식·가중치·업체수 판정·백테스트 엔진은 변경하지 않습니다.

'''
if section not in readme:
    anchor='## 중심모델군과 세부업종\n'
    if anchor not in readme:
        raise RuntimeError('README section anchor not found')
    readme=readme.replace(anchor,section+anchor,1)
README_PATH.write_text(readme, encoding='utf-8')

for test_path in (TEST_V157_PATH,TEST_V158_PATH):
    text=test_path.read_text(encoding='utf-8')
    text=text.replace('MODEL_VERSION = "v2.15.8"','MODEL_VERSION = "v2.15.9"')
    text=text.replace('투찰전략 분석 시스템 v2.15.8','투찰전략 분석 시스템 v2.15.9')
    text=text.replace('# 투찰전략 분석 시스템 v2.15.8','# 투찰전략 분석 시스템 v2.15.9')
    text=text.replace('values["모델버전"],"v2.15.8"','values["모델버전"],"v2.15.9"')
    test_path.write_text(text,encoding='utf-8')

TEST_V159_PATH.write_text('''"""Regression tests for v2.15.9 lower-limit parsing and summary Excel."""
import io
from pathlib import Path
import sys
import unittest

import numpy as np
from openpyxl import Workbook, load_workbook

sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from test_v2_15_4 import load_app_namespace


class StrategySummaryV2159Tests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app=load_app_namespace()

    def test_lower_limit_parser(self):
        parse=self.app["parse_lower_limit_rate"]
        self.assertEqual(parse(89.745),89.745)
        self.assertEqual(parse("89.745%"),89.745)
        self.assertEqual(parse(" 88 "),88.0)
        self.assertEqual(parse(0.89745),89.745)
        self.assertIsNone(parse(None))
        self.assertIsNone(parse(""))
        self.assertIsNone(parse("오류"))
        self.assertIsNone(parse(np.nan))

    def test_xlsx_parser_reads_lower_limit(self):
        wb=Workbook(); ws=wb.active
        ws.append([None,None,"샘플 공고"])
        ws.append(["분류","번호","공고명","공고번호","기초금액","투찰마감","지역","참가마감","투찰유무","발주기관","업종","낙찰하한율"])
        ws.append(["공사",1,"샘플 전기공사","B1",100000000,"26.09.30 (14:00)","전국","","미투찰","발주처","전기","89.745%"])
        buf=io.BytesIO(); wb.save(buf)
        bids=self.app["parse_xls"](buf.getvalue(),"sample.xlsx")
        self.assertEqual(len(bids),1)
        self.assertEqual(bids[0]["lower_limit_rate"],89.745)

    def test_both_parser_branches_capture_lower_limit(self):
        app_text=(Path(__file__).resolve().parents[1]/"입찰 앱.py").read_text(encoding="utf-8")
        self.assertEqual(app_text.count('"lower_limit_rate": parse_lower_limit_rate(row.get("낙찰하한율"))'),2)

    def test_summary_excel_layout_and_values(self):
        results=[
            {"bid":{"name":"1개사 공고","lower_limit_rate":89.745},"scope_info":{"company_count":1},"recommendations":[{"rate":0.0161}]},
            {"bid":{"name":"2개사 공고","lower_limit_rate":"88%"},"scope_info":{"company_count":2},"recommendations":[{"rate":-0.5479},{"rate":0.0164}]},
            {"bid":{"name":"3개사 공고","lower_limit_rate":None},"scope_info":{"company_count":3},"recommendations":[{"rate":-0.7154},{"rate":-0.1254},{"rate":-0.36}]},
        ]
        out=self.app["make_strategy_summary_excel"](results)
        wb=load_workbook(io.BytesIO(out.getvalue()))
        ws=wb["투찰전략"]
        self.assertEqual([ws.cell(2,c).value for c in range(1,6)],["공고명","낙찰하한율","업체1추천","업체2추천","업체3추천"])
        self.assertEqual([ws.cell(3,c).value for c in range(1,6)],["1개사 공고","89.745%","+0.0161%","참여대상 없음","참여대상 없음"])
        self.assertEqual([ws.cell(4,c).value for c in range(1,6)],["2개사 공고","88%","-0.5479%","+0.0164%","참여대상 없음"])
        self.assertEqual([ws.cell(5,c).value for c in range(1,6)],["3개사 공고","미확인","-0.7154%","-0.1254%","-0.3600%"])
        self.assertEqual(ws.freeze_panes,"A3")
        self.assertEqual(ws.auto_filter.ref,"A2:E5")

    def test_download_button_and_version_present(self):
        app_text=(Path(__file__).resolve().parents[1]/"입찰 앱.py").read_text(encoding="utf-8")
        self.assertIn('MODEL_VERSION = "v2.15.9"',app_text)
        self.assertIn("📥 투찰전략 요약 다운로드",app_text)
        self.assertIn("투찰전략_요약_",app_text)


if __name__=="__main__":
    unittest.main()
''',encoding='utf-8')

VERIFY_PATH.write_text('''# v2.15.9 구현·검증 결과

- 모델 버전 및 화면 상단 버전: v2.15.9
- XLS/XLSX 입찰서류함 `낙찰하한율` 파싱: 반영
- 요약 Excel 열: 공고명 → 낙찰하한율 → 업체1추천 → 업체2추천 → 업체3추천
- 1·2·3개사 참여대상 표시: 반영
- 기존 상세 Excel 및 화면 요약표: 유지
- 추천식·가중치·업체수 판정·백테스트 엔진: 변경 없음
- Python 문법 검사: GitHub Actions 통과 예정
- 핵심 단위·Excel 회귀검사: GitHub Actions 통과 예정
- v2.15.9 신규 파싱·Excel 검사: GitHub Actions 통과 예정
- Streamlit 부팅 및 첫 화면 HTTP 검사: GitHub Actions 통과 예정
- git diff --check: GitHub Actions 통과 예정
''',encoding='utf-8')
