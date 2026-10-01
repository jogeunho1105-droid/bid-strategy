import io
import json
import sys
import tempfile
import unittest
from unittest.mock import patch
from pathlib import Path
import numpy as np
import pandas as pd
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import bid_engine as engine
import bid_strategies as strategies
import ledger_metadata
from analysis.ledger_regression import live_adapter,cached_interval,ordered_window,cached_weights
from test_strategies_v2_15_11 import history,bid
from test_v2_15_4 import load_app_namespace

class CollectionTests(unittest.TestCase):
    def test_accelerated_full_strategy_loop_matches_real_app(self):
        modes=['S1','S2','S3','S4','S5','auto']
        for name,industry,org in [('일반 설계용역','설계','시험시청'),('건설사업관리 용역','건설사업관리','시험시청'),('전기공사','전기','시험시청'),('전력시설물 감리용역','감리','한국전력공사 부산울산본부')]:
            h=history(**{'공고명':name,'업종':industry,'발주기관':org});b=bid(name=name,industry=industry,org=org)
            expected={mode:strategies.recommend_bids([b],h,mode=mode)[0] for mode in modes}
            weights=cached_weights(strategies._weights)
            with patch.object(strategies,'_window',ordered_window),patch.object(strategies,'_weights',weights),patch.object(strategies,'interval_point',cached_interval()):
                adapter=live_adapter();e=engine.PredictionEngine()
                for r in engine.prepare_history(h).to_dict('records'):e.add(engine.source_record(r))
                row=engine.bid_record(b)
                for mode in modes:
                    actual=adapter(e,row,[b],mode,engine.load_policy(),'test')
                    self.assertEqual(actual['rates'],expected[mode]['rates'],(name,mode))
                    self.assertEqual(actual['selected_strategy'],expected[mode]['selected_strategy'])
    def test_chronological_window_and_weights_exact_parity(self):
        records=[engine.source_record(r) for r in engine.prepare_history(history(n=900)).to_dict('records')]
        weights=cached_weights(strategies._weights)
        for date in [pd.Timestamp('2026-06-01'),pd.Timestamp('2027-01-01'),pd.Timestamp('2029-01-01')]:
            for days,cap in [(90,5),(730,2000),(1460,300)]:
                expected=strategies._window(records,date,days,cap);actual=ordered_window(records,date,days,cap)
                self.assertEqual(actual,expected);np.testing.assert_array_equal(weights(actual,date),strategies._weights(expected,date))
    def test_cached_interval_bitwise_parity(self):
        d=engine.prepare_history(history(n=600));records=[engine.source_record(r) for r in d.to_dict('records')]
        point=cached_interval();date=pd.Timestamp('2028-01-01')
        for family,issuer in [(records,records[-80:]),(records,[]),(records[-20:],records[-5:])]:
            for retained in [[],[-.3],[.2,-.1]]:
                self.assertEqual(point(family,issuer,date,retained,.01),strategies.interval_point(family,issuer,date,retained,.01))
    def test_streaming_adapter_matches_live_all_modes(self):
        h=history();b=bid();row=engine.bid_record(b);e=engine.PredictionEngine()
        for r in engine.prepare_history(h).to_dict('records'):e.add(engine.source_record(r))
        adapter=live_adapter()
        for mode in ['S1','S2','S3','S4','S5','auto']:
            expected=strategies.recommend_bids([b],h,mode=mode)[0]
            actual=adapter(e,row,[b],mode,engine.load_policy(),'test')
            for key in ['rates','selected_strategy','family_n','winner_n','latest_train_date','fallback_notes']:
                self.assertEqual(actual[key],expected[key],(mode,key))
    def test_new_metadata_cutoff_preserves_missing_and_future(self):
        app=load_app_namespace();h=history();h.loc[0,engine.WINNER]=np.nan
        h.loc[1,'개찰일']=pd.Timestamp('2027-01-01')
        h.attrs[ledger_metadata.ATTR]={'as_of':'2026-10-01'}
        d=app['prepare_history_frame'](h)
        self.assertTrue(d.loc[d['공고번호'].eq('H00000-1'),engine.WINNER].isna().all())
        self.assertFalse(d['공고번호'].eq('H00001-1').any())
        self.assertEqual(len(h),90)
    def test_metadata_sheet_preserves_all_existing_cells(self):
        from openpyxl import Workbook,load_workbook
        wb=Workbook();ws=wb.active;ws.title='투찰전략';ws.append(['공고명','낙찰하한율']);ws.append(['시험',88.])
        b=io.BytesIO();wb.save(b);b.seek(0)
        m=dict(source='수집원본',version='test',source_rows=10,operational_rows=9,training_rows=8,data_end_date='2026-09-29',completed_tasks=251,total_tasks=252,as_of='2026-10-01',region_basis='참가제한 별도 확인',unresolved='미해결1건',numeric_status={'both_numeric':5,'target_only_numeric':2,'target_not_numeric':3},warnings={k:0 for k in ['missing_name','missing_org','missing_date','future_opening','abnormal_target','abnormal_winner']},dual_profile_events=1)
        out=ledger_metadata.add_excel_metadata(b,m);result=load_workbook(out)
        self.assertEqual(list(result['투찰전략'].values),list(ws.values));self.assertIn('데이터출처',result.sheetnames)
    def test_upload_rejects_false_row_count(self):
        from openpyxl import Workbook
        wb=Workbook();wb.active.append(['공고명']);wb.active.append(['시험'])
        meta=wb.create_sheet('데이터출처');meta.append(['출처']);meta.append([json.dumps({'operational_rows':100})])
        buf=io.BytesIO();wb.save(buf)
        with self.assertRaises(ValueError):ledger_metadata.read_upload(buf.getvalue())

if __name__=='__main__':unittest.main()
