from __future__ import annotations

import ast
import io
import json
import os
import re
import unittest
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook


APP = Path(__file__).resolve().parents[1] / "입찰 앱.py"


def load_app_namespace() -> dict:
    tree = ast.parse(APP.read_text(encoding="utf-8"))
    namespace = {
        "pd": pd,
        "np": np,
        "re": re,
        "io": io,
        "os": os,
        "json": json,
        "datetime": datetime,
    }
    bodies = []
    for node in tree.body:
        if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)):
            node.decorator_list = []
            bodies.append(node)
        elif isinstance(node, ast.Assign):
            try:
                value = ast.literal_eval(node.value)
            except Exception:
                continue
            for target in node.targets:
                if isinstance(target, ast.Name):
                    namespace[target.id] = value
    module = ast.Module(body=bodies, type_ignores=[])
    ast.fix_missing_locations(module)
    exec(compile(module, str(APP), "exec"), namespace)
    return namespace


class V2154Tests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = load_app_namespace()

    def history(self, org="일반기관", name="일반 설계용역", industry="설계"):
        dates = pd.date_range("2026-06-01", periods=12, freq="7D")
        values = np.linspace(0.10, 0.21, len(dates))
        return pd.DataFrame(
            {
                "개찰일": dates,
                "공고번호": [f"N{i:02d}" for i in range(len(dates))],
                "번호": range(len(dates)),
                "발주기관": org,
                "공고명": name,
                "업종": industry,
                "기초금액": 100_000_000,
                "업체수": 20,
                "지역": "부산",
                "예가/기초(0%)": values,
            }
        )

    def test_stable_sort_and_exact_duplicate_removal(self):
        frame = pd.DataFrame(
            {
                "개찰일": ["2026-01-01"] * 4,
                "공고번호": ["B", "A", "C", "A"],
                "번호": [2, 1, 2, 1],
                "예가/기초(0%)": [0.2, 0.1, 0.3, 0.1],
            }
        )
        ordered = self.app["history_sorted"](frame)
        self.assertEqual(ordered["공고번호"].tolist(), ["A", "B", "C"])
        self.assertEqual(len(ordered), 3)

    def test_general_overlay_uses_v2154_weights(self):
        frame = self.app["enrich_history"](self.history())
        bid = {
            "org": "일반기관",
            "name": "일반 설계용역",
            "industry": "설계",
            "base": 100_000_000,
            "deadline": datetime(2026, 8, 31),
        }
        _, note = self.app["recent_volatility_overlay"](bid, frame, 0.0)
        self.assertIn("v2.15.4 보수화", note)
        self.assertIn("최근가중 0.10", note)
        self.assertIn("×0.04", note)

    def test_electric_overlay_keeps_v2153_weights(self):
        frame = self.app["enrich_history"](
            self.history(name="배전선로 전기공사", industry="전기")
        )
        bid = {
            "org": "일반기관",
            "name": "배전선로 전기공사",
            "industry": "전기",
            "base": 100_000_000,
            "deadline": datetime(2026, 8, 31),
        }
        _, note = self.app["recent_volatility_overlay"](bid, frame, 0.0)
        self.assertIn("v2.15.3 유지", note)
        self.assertIn("최근가중 0.18", note)
        self.assertIn("×0.08", note)

    def test_kepco_local_single_keeps_v2153_weights(self):
        org = "한국전력공사 경북본부"
        frame = self.app["enrich_history"](
            self.history(org=org, name="배전공사 감리용역", industry="전력감리")
        )
        bid = {
            "org": org,
            "name": "배전공사 감리용역",
            "industry": "전력감리",
            "base": 100_000_000,
            "deadline": datetime(2026, 8, 31),
        }
        _, note = self.app["recent_volatility_overlay"](bid, frame, 0.0)
        self.assertIn("v2.15.3 유지", note)
        self.assertIn("최근가중 0.18", note)

    def test_quality_summary_and_reference_amount(self):
        frame = pd.DataFrame(
            [
                {"발주기관": "A", "공고명": "전기공사", "업종": "전기", "지역": "", "예가/기초(0%)": 0.1},
                {"발주기관": "A", "공고명": "전기공사", "업종": "전기", "지역": "", "예가/기초(0%)": 0.1},
                {"발주기관": "B", "공고명": "설계용역", "업종": "설계", "지역": None, "예가/기초(0%)": None},
            ]
        )
        quality = self.app["history_quality_summary"](frame)
        self.assertEqual(quality["exact_duplicates"], 1)
        self.assertEqual(quality["invalid_target_rows"], 1)
        self.assertEqual(quality["electric_region_missing_rows"], 1)
        self.assertEqual(
            self.app["recommendation_reference_amount"](100_000_000, 0.1),
            100_100_000,
        )

    def test_excel_contains_v2154_audit_columns(self):
        results = [
            {
                "bid": {
                    "no": 1,
                    "bid_no": "N-1",
                    "name": "일반 설계용역",
                    "org": "일반기관",
                    "industry": "설계",
                    "region": "부산",
                    "base": 100_000_000,
                    "deadline": "2026-08-31",
                },
                "scope_info": {"applicable": False, "scope": "기존분석", "company_count": 3},
                "recommendations": [
                    {
                        "rate": 0.1,
                        "model_label": "중심모델",
                        "basis_n": 12,
                        "basis": "최근90일 오버레이(v2.15.4 보수화, 기준표본 n=12, 최근 n=10, 변동성 0.2000, 최근가중 0.10)",
                    }
                ],
            }
        ]
        quality = {
            "exact_duplicates": 0,
            "invalid_target_rows": 0,
            "region_missing_rows": 0,
            "region_missing_rate": 0.0,
            "electric_region_missing_rows": 0,
            "electric_rows": 0,
            "electric_region_missing_rate": 0.0,
        }
        output = self.app["make_excel_simple"](results, quality)
        wb = load_workbook(io.BytesIO(output.getvalue()), data_only=False)
        ws = wb["업체별 추천"]
        headers = [cell.value for cell in ws[2]]
        self.assertIn("업체1 추천기준금액(원)", headers)
        self.assertIn("업체1 최근90일가중치", headers)
        self.assertIn("데이터품질경고", headers)
        self.assertEqual(ws.cell(3, 11).value, 100_100_000)
        self.assertEqual(ws.cell(3, len(headers)).value, "v2.15.6")
        self.assertEqual(wb.sheetnames, ["업체별 추천", "사후낙찰검증", "검증기준"])
        audit_ws = wb["사후낙찰검증"]
        self.assertEqual(audit_ws.cell(3, 1).value, "N-1")
        self.assertEqual(audit_ws.cell(3, 6).value, 0.1)
        self.assertEqual(audit_ws.cell(3, 11).value, '=IFERROR(CHOOSE(J3,F3,G3,H3),"")')
        self.assertIn("COUNT(K3,M3,N3)<3", audit_ws.cell(3, 18).value)
        self.assertIn("AND(M3<K3,K3<N3)", audit_ws.cell(3, 18).value)
        self.assertIn("COUNT(L3,M3,N3)<3", audit_ws.cell(3, 19).value)
        self.assertIn("AND(M3<L3,L3<N3)", audit_ws.cell(3, 19).value)
        basis_ws = wb["검증기준"]
        self.assertEqual(basis_ws.cell(2, 2).value, "2026-09-01")
        self.assertEqual(basis_ws.cell(6, 2).value, 0.1423506011)

    def test_quality_summary_includes_history_period(self):
        frame = self.history()
        quality = self.app["history_quality_summary"](frame)
        self.assertEqual(quality["data_start_date"], "2026-06-01")
        self.assertEqual(quality["data_end_date"], "2026-08-17")
        self.assertEqual(
            self.app["history_period_text"](quality),
            "이력 기간 2026-06-01 ~ 2026-08-17",
        )

    def test_busan_local_company_count_amount_boundaries(self):
        common = {
            "org": "한국전력공사 부산울산본부",
            "name": "부산 지역제한 감리용역",
            "industry": "전력감리",
            "deadline": datetime(2026, 8, 31),
        }
        cases = [
            (99_999_999, "지역제한", 2),
            (100_000_000, "지역제한", 3),
            (252_999_999, "지역제한", 3),
            (253_000_000, "전국입찰", 3),
        ]
        for base, expected_scope, expected_count in cases:
            with self.subTest(base=base):
                scope = self.app["classify_kepco_scope"]({**common, "base": base})
                self.assertEqual(scope["scope"], expected_scope)
                self.assertEqual(scope["company_count"], expected_count)

        below = self.app["classify_kepco_scope"]({**common, "base": 99_999_999})
        at_threshold = self.app["classify_kepco_scope"]({**common, "base": 100_000_000})
        self.assertIn("기초금액 1억원 미만", below["basis"])
        self.assertIn("기초금액 1억원 이상", at_threshold["basis"])

    def test_historical_notice_amount_boundaries_remain_unchanged(self):
        common = {
            "org": "한국전력공사 부산울산본부",
            "name": "부산 지역제한 감리용역",
            "industry": "전력감리",
        }
        cases = [
            (datetime(2022, 12, 31), 230_999_999, "지역제한", 3),
            (datetime(2022, 12, 31), 231_000_000, "전국입찰", 3),
            (datetime(2024, 12, 31), 241_999_999, "지역제한", 3),
            (datetime(2024, 12, 31), 242_000_000, "전국입찰", 3),
        ]
        for deadline, base, expected_scope, expected_count in cases:
            with self.subTest(deadline=deadline, base=base):
                scope = self.app["classify_kepco_scope"](
                    {**common, "deadline": deadline, "base": base}
                )
                self.assertEqual(scope["scope"], expected_scope)
                self.assertEqual(scope["company_count"], expected_count)

    def test_existing_kepco_and_non_supervision_counts_do_not_change(self):
        deadline = datetime(2026, 8, 31)
        for org in ("한국전력공사 경북본부", "한국전력공사 대구본부"):
            with self.subTest(org=org):
                scope = self.app["classify_kepco_scope"](
                    {
                        "org": org,
                        "name": "배전공사 감리용역",
                        "industry": "전력감리",
                        "base": 100_000_000,
                        "deadline": deadline,
                    }
                )
                self.assertEqual(scope["scope"], "지역제한")
                self.assertEqual(scope["company_count"], 1)

        non_supervision = self.app["classify_kepco_scope"](
            {
                "org": "한국전력공사 부산울산본부",
                "name": "배전선로 설계용역",
                "industry": "전력설계",
                "base": 100_000_000,
                "deadline": deadline,
            }
        )
        self.assertFalse(non_supervision["applicable"])
        self.assertEqual(non_supervision["company_count"], 3)

    def test_busan_three_company_rule_exposes_line_hedge(self):
        recommendations = [
            {"company": "업체 1", "rate": -0.1, "role": "방향성 헷지"},
            {"company": "업체 2", "rate": 0.0, "role": "중심모델"},
            {"company": "업체 3", "rate": 0.1, "role": "라인 헷지"},
        ]
        bid = {
            "org": "한국전력공사 부산울산본부",
            "name": "부산 지역제한 감리용역",
            "industry": "전력감리",
            "base": 100_000_000,
            "deadline": datetime(2026, 8, 31),
        }
        scope = self.app["classify_kepco_scope"](bid)
        selected = self.app["apply_company_count"](recommendations, scope, None, bid)
        self.assertEqual(len(selected), 3)
        self.assertEqual(selected[2]["role"], "라인 헷지")

    def test_busan_three_company_excel_contains_company3(self):
        bid = {
            "no": 1,
            "bid_no": "BUSAN-100M",
            "name": "부산 지역제한 감리용역",
            "org": "한국전력공사 부산울산본부",
            "industry": "전력감리",
            "region": "부산",
            "base": 100_000_000,
            "deadline": "2026-08-31",
        }
        scope = self.app["classify_kepco_scope"](bid)
        recommendations = [
            {
                "company": f"업체 {position}",
                "rate": rate,
                "role": role,
                "model_label": role,
                "basis_n": 12,
                "basis": f"{role} 검증",
            }
            for position, rate, role in (
                (1, -0.1, "방향성 헷지"),
                (2, 0.0, "중심모델"),
                (3, 0.1, "라인 헷지"),
            )
        ]
        output = self.app["make_excel_simple"](
            [{"bid": bid, "scope_info": scope, "recommendations": recommendations}],
            {},
        )
        wb = load_workbook(io.BytesIO(output.getvalue()), data_only=False)
        ws = wb["업체별 추천"]
        headers = [cell.value for cell in ws[2]]
        row = {header: ws.cell(3, index + 1).value for index, header in enumerate(headers)}
        self.assertEqual(row["참여업체수"], 3)
        self.assertEqual(row["업체3 추천사정률(%)"], 0.1)
        self.assertEqual(row["업체3 추천기준금액(원)"], 100_100_000)
        self.assertEqual(row["업체3 적용모델"], "라인 헷지")
        self.assertIn("입찰구분: 지역제한", row["입찰판정근거"])
        self.assertIn("참여업체: 3개사", row["입찰판정근거"])
        self.assertIn("기초금액 1억원 이상", row["입찰판정근거"])
        self.assertEqual(row["모델버전"], "v2.15.6")


if __name__ == "__main__":
    unittest.main()
