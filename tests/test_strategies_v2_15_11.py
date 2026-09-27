import copy
import io
import json
from pathlib import Path
import sys
import unittest
from unittest.mock import patch
import numpy as np
import pandas as pd
from openpyxl import Workbook,load_workbook
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import bid_strategies as strategy
import bid_engine as engine
from test_v2_15_4 import load_app_namespace


def bid(**changes):
    b=dict(no=1,bid_no="T2026001-1",name="일반 설계용역",industry="설계",org="시험시청",base=150_000_000,base_억=1.5,deadline="2026-09-30",region="전국",lower_limit_rate=88.)
    b.update(changes);return b


def history(n=90,**changes):
    vals=np.sin(np.arange(n)*.83)*.45
    d=pd.DataFrame({"개찰일":pd.date_range("2026-06-01",periods=n),"공고번호":[f"H{i:05d}-1" for i in range(n)],"번호":range(n),"공고명":["일반 설계용역"]*n,"발주기관":["시험시청"]*n,"업종":["설계"]*n,"지역":["전국"]*n,"기초금액":[150_000_000]*n,"예가/기초(0%)":vals,"1순위사정율(0%)":vals+.11})
    for k,v in changes.items():d[k]=v
    return d


class RevisionTests(unittest.TestCase):
    def test_numeric_revision_not_lexical(self):
        current,removed=strategy.latest_notices([bid(bid_no=f"E1-{x}") for x in (1,2,10)])
        self.assertEqual([b["bid_no"] for b in current],["E1-10"]);self.assertEqual(len(removed),2)
    def test_latest_rate_is_retained(self):
        b=[bid(bid_no="E1-1",lower_limit_rate=88.745),bid(bid_no="E1-2",lower_limit_rate=89.745)]
        kept,_=strategy.latest_notices(b)
        self.assertEqual(kept[0]["lower_limit_rate"],89.745);self.assertEqual(len(b),2)
    def test_fifteen_to_thirteen(self):
        b=[bid(bid_no=f"E{i}-2",no=i) for i in range(13)]
        kept,removed=strategy.latest_notices(b+[bid(bid_no="E0-1"),bid(bid_no="E1-1")])
        self.assertEqual(len(kept),13);self.assertEqual(len(removed),2)
    def test_missing_ids_and_separate_lots_not_merged(self):
        kept,_=strategy.latest_notices([bid(bid_no=""),bid(bid_no=""),bid(lot_no="A"),bid(lot_no="B")]);self.assertEqual(len(kept),4)
    def test_same_highest_conflict_is_blocked(self):
        kept,_=strategy.latest_notices([bid(),bid(base=99_000_000)])
        self.assertTrue(all("revision_warning" in v for v in kept))
        self.assertTrue(all(x["status"]=="산정보류" for x in strategy.recommend_bids(kept,history(),mode="S3")))
    def test_duplicate_highest_identical(self):
        kept,removed=strategy.latest_notices([bid(no=1),bid(no=2)]);self.assertEqual((len(kept),len(removed)),(1,1))
    def test_cancellation_does_not_revive_old(self):
        kept,_=strategy.latest_notices([bid(bid_no="E1-1"),bid(bid_no="E1-2",name="일반 설계용역 [취소]")])
        self.assertEqual(kept[0]["bid_no"],"E1-2")
        self.assertEqual(strategy.recommend_bids(kept,history(),mode="S3")[0]["rates"],[])


class StrategyTests(unittest.TestCase):
    def test_auto_routes_and_no_retroactive_policy(self):
        self.assertEqual(strategy.resolved_strategy("electric_construction","auto","2026-09-30"),"S4")
        self.assertEqual(strategy.resolved_strategy("construction_management","auto","2026-09-30"),"S3")
        self.assertEqual(strategy.resolved_strategy("design","auto","2026-09-30"),"S3")
        self.assertEqual(strategy.resolved_strategy("design","auto","2026-09-27"),"S1")
    def test_s1_equals_real_shared_engine(self):
        b=bid();h=history();expected=engine.recommend_bid(b,h,{"rules":{}})
        actual=strategy.recommend_bids([b],h,{"rules":{}},"S1")[0];self.assertEqual(actual["rates"],expected["rates"])
    def test_s2_and_s4_equal_shared_engine_sufficient_pool(self):
        b=bid();h=history();expected=engine.recommend_bid(b,h,{"rules":{}})
        for s,es in [("S2","family_center"),("S4","interval_center")]:
            p=strategy.recommend_bids([b],h,{"rules":{}},s)[0];self.assertEqual(p["rates"],expected["recommendations"][es])
    def test_s3_matches_independent_weighted_cdf(self):
        p=strategy.recommend_bids([bid()],history(),mode="S3")[0];h=history();vals=h[engine.RATE].to_numpy()
        w=np.power(2.,-(pd.Timestamp("2026-09-30")-h["개찰일"]).dt.days.to_numpy()/180.)
        ix=np.argsort(vals);cw=np.cumsum(w[ix]);expected=[round(float(vals[ix[min(np.searchsorted(cw,q*cw[-1]),len(vals)-1)]]),4) for q in (.2,.5,.8)]
        self.assertEqual(p["rates"],expected)
    def test_same_day_and_future_outcomes_not_used(self):
        h=history();extra=h.tail(2).copy();extra["개찰일"]=["2026-09-30","2026-10-01"];extra["공고번호"]=["F1","F2"];extra[engine.RATE]=8.;extra[engine.WINNER]=9.
        for s in strategy.LABELS:
            x=strategy.recommend_bids([bid()],h,mode=s)[0]
            y=strategy.recommend_bids([bid()],pd.concat([h,extra],ignore_index=True),mode=s)[0]
            self.assertEqual(x["rates"],y["rates"]);self.assertLess(y["latest_train_date"],"2026-09-30")
    def test_single_company_s4_s5_equal(self):
        h=history(**{"공고명":"전기공사","업종":"전기"});b=bid(name="전기공사",industry="전기")
        p=[strategy.recommend_bids([b],h,mode=s)[0] for s in ("S4","S5")]
        self.assertEqual(len(p[0]["rates"]),1);self.assertEqual(p[0]["rates"],p[1]["rates"])
    def test_1_2_3_company_counts_same_for_all_strategies(self):
        cases=[("설계","설계","시험시청",150_000_000,3),("배전 감리","전력감리","한국전력공사 부산울산본부",99_999_999,2),("배전 감리","전력감리","한국전력공사 부산울산본부",100_000_000,3),("배전 감리","전력감리","한국전력공사 대구본부",80_000_000,1)]
        for name,ind,org,base,count in cases:
            b=bid(name=name,industry=ind,org=org,base=base);h=history(**{"공고명":name,"업종":ind,"발주기관":org,"기초금액":base})
            for s in strategy.LABELS:self.assertEqual(len(strategy.recommend_bids([b],h,mode=s)[0]["rates"]),count)
    def test_missing_winner_falls_back_not_fake_interval_probability(self):
        h=history();h[engine.WINNER]=np.nan;p=strategy.recommend_bids([bid()],h,mode="S4")[0]
        self.assertEqual(p["selected_strategy"],"S3");self.assertEqual(p["winner_n"],0);self.assertIn("승계 안 함"," ".join(p["fallback_notes"]))
    def test_empty_history_and_invalid_date_hold(self):
        self.assertEqual(strategy.recommend_bids([bid()],None)[0]["rates"],[])
        self.assertEqual(strategy.recommend_bids([bid(deadline="oops")],history())[0]["rates"],[])
    def test_tiny_pool_reference_and_hold(self):
        self.assertEqual(strategy.recommend_bids([bid()],history(4),mode="S3")[0]["rates"],[])
        p=strategy.recommend_bids([bid()],history(10),mode="S3")[0]
        self.assertEqual(p["selected_strategy"],"REFERENCE");self.assertEqual(len(set(p["rates"])),1)
    def test_stale_but_valid_data_expands_with_warning(self):
        h=history();h["개찰일"]-=pd.Timedelta(days=800);p=strategy.recommend_bids([bid()],h,mode="S3")[0]
        self.assertEqual(p["window_days"],1460);self.assertEqual(p["recent90_n"],0);self.assertTrue(p["fallback_notes"])
    def test_no_cross_sector_fallback(self):
        self.assertEqual(strategy.recommend_bids([bid(name="전기공사",industry="전기")],history(),mode="S3")[0]["rates"],[])
    def test_local_and_national_pools_not_mixed(self):
        b=bid(name="배전 감리",industry="전력감리",org="한국전력공사 부산울산본부",base=100_000_000)
        h=history(**{"공고명":"배전 감리","업종":"전력감리","발주기관":b["org"],"기초금액":100_000_000})
        x=strategy.recommend_bids([b],h,mode="S3")[0]
        national=h.copy();national["공고번호"]="N"+national["공고번호"];national["기초금액"]=500_000_000;national[engine.RATE]=3.
        y=strategy.recommend_bids([b],pd.concat([h,national],ignore_index=True),mode="S3")[0]
        self.assertEqual(x["rates"],y["rates"])
    def test_unit_supervision_not_classified_by_unit_amount(self):
        b=bid(name="2027년 배전공사 감리용역 단가",industry="전력감리",org="한국전력공사 부산울산본부",base=10_000_000)
        p=strategy.recommend_bids([b],history(),mode="S3")[0]
        self.assertEqual(p["rates"],[]);self.assertIn("총 추정가격",p["error"])
    def test_inputs_not_mutated_and_snapshot_serializable(self):
        b=bid();before=copy.deepcopy(b);h=history();previous=h.copy(deep=True);p=strategy.recommend_bids([b],h,mode="S3")[0]
        self.assertEqual(json.loads(strategy.snapshot([{"bid":b,"prediction":p}]))["model_version"],"v2.15.11")
        self.assertEqual(b,before);pd.testing.assert_frame_equal(h,previous)


class AppIntegrationTests(unittest.TestCase):
    def test_application_adapter_and_both_excels(self):
        app=load_app_namespace();b=bid();h=history();result=app["predict_final_batch"]([b],h,strategy_mode="S3")[0]
        result.update(bid=b,scope_info=app["classify_bid_scope"](b));p=result["prediction"]
        self.assertEqual([r["rate"] for r in result["recommendations"]],p["rates"])
        for func in ("make_excel_simple","make_strategy_summary_excel"):
            base=app[func]([result]);out=strategy.add_review_sheets(base,[result],[{"제외공고번호":"OLD-1","적용공고번호":"OLD-2"}]);wb=load_workbook(out)
            self.assertIn("전략선택근거",wb.sheetnames);self.assertIn("정정공고제외",wb.sheetnames)
            self.assertEqual(wb["전략선택근거"]["D2"].value,"S3");self.assertEqual(wb["전략선택근거"]["O2"].value,"v2.15.11")
    def test_first_screen_app_execution(self):
        from streamlit.testing.v1 import AppTest
        root=Path(__file__).resolve().parents[1];at=AppTest.from_file(str(root/"입찰 앱.py"),default_timeout=40).run()
        self.assertEqual(len(at.exception),0);self.assertTrue(any("v2.15.11" in m.value for m in at.markdown));self.assertTrue(any(s.label=="전략 선택" for s in at.selectbox))
    def test_app_with_synthetic_upload_exercises_real_ui(self):
        from streamlit.testing.v1 import AppTest
        import streamlit as st
        root=Path(__file__).resolve().parents[1];wb=Workbook();ws=wb.active;ws.append(["sample"])
        ws.append(["번호","공고명","공고번호","기초금액","투찰마감","발주기관","업종","지역","낙찰하한율"])
        for i,rev in enumerate((1,2),1):ws.append([i,"일반 설계용역",f"T2026001-{rev}",150_000_000,"26.09.30 (10:00)","시험시청","설계","전국",88.])
        upload=io.BytesIO();wb.save(upload);upload.seek(0);upload.name="sample.xlsx"
        path=root/"data"/"history.pkl";path.parent.mkdir(exist_ok=True);previous=path.read_bytes() if path.exists() else None
        history().to_pickle(path);st.cache_data.clear()
        try:
            with patch.object(st,"file_uploader",return_value=upload):at=AppTest.from_file(str(root/"입찰 앱.py"),default_timeout=60).run()
            self.assertEqual(len(at.exception),0,repr(at.exception));self.assertTrue(any("최종 1건" in s.value for s in at.success))
            self.assertTrue(any("정정공고 제외내역" in x.label for x in at.expander));self.assertTrue(len(at.dataframe)>=1)
        finally:
            if previous is None:path.unlink(missing_ok=True)
            else:path.write_bytes(previous)
            st.cache_data.clear()


if __name__=="__main__":unittest.main()
