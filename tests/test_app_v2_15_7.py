"""Verify the app displays the shared-engine outputs without another correction."""
import io
from pathlib import Path
import sys
import unittest
from unittest.mock import patch

import numpy as np
import pandas as pd
from openpyxl import load_workbook

sys.path.insert(0,str(Path(__file__).resolve().parents[1]))

import bid_engine
from test_v2_15_4 import load_app_namespace


class AppSharedEngineTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app=load_app_namespace()

    def history(self, name="VLF 진단", industry="진단", org="한국전력공사 부산울산본부"):
        n=45
        return pd.DataFrame({
            "개찰일":pd.date_range("2026-01-01",periods=n),
            "공고명":[name]*n,"업종":[industry]*n,"발주기관":[org]*n,
            "공고번호":[f"H{i:03d}" for i in range(n)],"기초금액":[80_000_000]*n,
            "예가/기초(0%)":np.sin(np.arange(n))*0.5,"1순위사정율(0%)":np.sin(np.arange(n))*0.5+0.05,
        })

    def bid(self, name="VLF 진단", industry="진단", org="한국전력공사 부산울산본부", base=80_000_000):
        return {"no":1,"bid_no":"B1","name":name,"industry":industry,"org":org,
                "base":base,"base_억":base/1e8,"deadline":"26.09.05","region":""}

    def test_final_rates_equal_shared_engine_no_double_correction(self):
        for name,industry,org,count in [
            ("VLF 진단","진단","한국전력공사 부산울산본부",2),
            ("배전 감리","전력감리","한국전력공사 경북본부",1),
            ("전기공사","전기","시청",1),
            ("일반 설계","설계","시청",3),
        ]:
            with self.subTest(name=name):
                history=self.history(name,industry,org)
                bid=self.bid(name,industry,org)
                expected=bid_engine.recommend_bid(bid,history,{"rules":{}})
                actual=self.app["predict_final_batch"]([bid],history,{"rules":{}})[0]
                self.assertEqual([r["rate"] for r in actual["recommendations"]],expected["rates"])
                self.assertEqual(len(actual["recommendations"]),count)

    def test_batch_uses_one_engine_call_and_handles_invalid_date(self):
        bids=[self.bid(),{**self.bid(),"deadline":"invalid"},self.bid(base=100_000_000)]
        with patch.object(bid_engine,"recommend_bids",wraps=bid_engine.recommend_bids) as shared:
            result=self.app["predict_final_batch"](bids,self.history(),{"rules":{}})
        self.assertEqual(shared.call_count,1)
        self.assertEqual(len(result[0]["recommendations"]),2)
        self.assertEqual(result[1]["recommendations"],[])
        self.assertIn("투찰마감일",result[1]["error"])
        self.assertEqual(len(result[2]["recommendations"]),3)

    def test_current_future_invalid_history_excluded_from_charts_and_rates(self):
        bid=self.bid()
        original=self.history()
        future=original.iloc[:3].copy()
        future["개찰일"]=["26.09.05","26.09.06","unknown"]
        future["예가/기초(0%)"]=[8.0,7.0,6.0]
        future["공고번호"]=["F1","F2","F3"]
        combined=pd.concat([original,future],ignore_index=True)
        baseline=self.app["predict_final_batch"]([bid],original,{"rules":{}})[0]
        result=self.app["predict_final_batch"]([bid],combined,{"rules":{}})[0]
        self.assertEqual(result["prediction"]["rates"],baseline["prediction"]["rates"])
        chart_history=self.app["history_for_bid"](bid,combined)
        self.assertEqual(len(chart_history),len(original))
        chart=self.app["simple_flow_data"](combined,bid)
        self.assertLess(max(chart["org_vals"]),1)

    def test_family_and_alternative_metadata_and_vlf_excel_slots(self):
        bid=self.bid()
        history=self.history()
        policy={"rules":{"diagnosis|기존분석":"interval_direction"}}
        row=self.app["predict_final_batch"]([bid],history,policy)[0]
        self.assertTrue(row["recommendations"][0]["alternative"])
        self.assertIn("대체",row["recommendations"][0]["role"])
        self.assertEqual(row["recommendations"][1]["role"],"중심모델")
        row.update(bid=bid,scope_info=self.app["classify_bid_scope"](bid))
        workbook=load_workbook(io.BytesIO(self.app["make_excel_simple"]([row],{}).getvalue()))
        ws=workbook["업체별 추천"]
        headers=[c.value for c in ws[2]]
        values={key:ws.cell(3,i+1).value for i,key in enumerate(headers)}
        self.assertEqual(values["중심모델군"],"일반진단")
        self.assertEqual(values["용역분류"],"VLF")
        self.assertEqual(values["업체3 추천사정률(%)"],"참여대상 없음")
        self.assertEqual(values["모델버전"],"v2.15.9")
        self.assertEqual(workbook["사후낙찰검증"].cell(3,8).value,None)
        self.assertIn("AND(M3<K3,K3<N3)",workbook["사후낙찰검증"].cell(3,18).value)
        self.assertNotIn("N3<=M3",workbook["사후낙찰검증"].cell(3,18).value)
        self.assertIn("COUNT(K3,M3,N3)<3",workbook["사후낙찰검증"].cell(3,18).value)
        self.assertIsNone(workbook["사후낙찰검증"].cell(3,14).comment)
        self.assertEqual(len(workbook["사후낙찰검증"].conditional_formatting),1)

    def test_summary_uses_dated_policy_without_stale_center_metric(self):
        policy={"version":"v2.15.7","as_of":"2026-09-04","effective_from":"2026-09-05",
                "summary":{"recent1y":{"공고수":100,"판정가능":90,"가상낙찰":9,"가상낙찰률":0.1},
                           "accepted_route_count":2}}
        with patch.object(bid_engine,"load_policy",return_value=policy):
            summary=self.app["load_audit_summary"]()
        self.assertEqual(summary["as_of"],"2026-09-04")
        self.assertEqual(summary["recent_1y_virtual_win_rate"],0.1)
        self.assertNotIn("recent_1y_center_or_single_rate",summary)
        with patch.object(bid_engine,"load_policy",return_value={"rules":{}}):
            previous=self.app["load_audit_summary"]()
        self.assertTrue(previous["previous_version"])
        self.assertEqual(previous["version"],"v2.15.6")

    def test_shared_cleanup_compatible_with_numeric_ui(self):
        frame=self.history()
        frame["예가/기초(0%)"]=frame["예가/기초(0%)"].astype(object)
        frame["개찰일"]=frame["개찰일"].astype(object)
        frame.loc[0,"예가/기초(0%)"]="0.1234"
        frame.loc[1,"개찰일"]="unknown"
        prepared=self.app["prepare_history_frame"](frame)
        self.assertEqual(len(prepared),len(frame)-1)
        self.assertTrue(pd.api.types.is_numeric_dtype(prepared["예가/기초(0%)"]))
        self.assertIn("_date",prepared)


if __name__=="__main__":
    unittest.main()
