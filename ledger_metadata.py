"""Data provenance, independent of recommendation formulas and scope rules."""
import hashlib
import io
import json
from pathlib import Path
import pandas as pd

ATTR = "collection_manifest"

def data_signature(frame, columns):
    import bid_engine
    data=frame[columns].copy()
    numeric_columns=['번호','기초금액','예정가격','1순위투찰금액','1순위기초대비','1순위사정율(100%)','1순위사정율(0%)','예가/기초(100%)','예가/기초(0%)','업체수']
    for col in columns:
        if col=='개찰일':data[col]=bid_engine.rules.parse_date_series(data[col]).dt.strftime('%Y-%m-%d').fillna('')
        elif col in numeric_columns:
            data[col]=bid_engine.numeric(data[col]).map(lambda v:format(v,'.12g') if pd.notna(v) else '')
        else:data[col]=data[col].fillna('').astype(str)
    return hashlib.sha256(data.to_csv(index=False,lineterminator='\n').encode('utf-8')).hexdigest()

def read_upload(content):
    book=pd.ExcelFile(io.BytesIO(content))
    frame=pd.read_excel(book,sheet_name=book.sheet_names[0],dtype={'공고번호':str,'1순위사업자번호':str})
    if "데이터출처" in book.sheet_names:
        meta=pd.read_excel(book,sheet_name="데이터출처",header=None)
        manifest=json.loads(meta.iloc[1,0])
        if len(frame)!=manifest["operational_rows"]:
            raise ValueError("신규 원장 행수와 출처정보가 일치하지 않습니다.")
        if data_signature(frame,manifest['original_columns'])!=manifest['operational_data_sha256']:
            raise ValueError("신규 원장 데이터와 검증 버전의 내용 해시가 일치하지 않습니다.")
        frame.attrs[ATTR]=manifest
        if '원천행출처' not in book.sheet_names:
            raise ValueError('전체 원천행 출처 시트가 없습니다.')
        lineage=pd.read_excel(book,sheet_name='원천행출처')
        if len(lineage)!=manifest['source_rows']:
            raise ValueError('원천행 출처 행수가 검증 버전과 다릅니다.')
        import gzip
        frame.attrs['collection_lineage_csv_gzip']=gzip.compress(lineage.to_csv(index=False,lineterminator='\n').encode('utf-8-sig'))
        if manifest.get('classification_review_events'):
            if '분류확인' not in book.sheet_names:raise ValueError('두 출처 분류확인 시트가 없습니다.')
            review=pd.read_excel(book,sheet_name='분류확인',dtype={'공고번호':str})
            if len(review)!=manifest['classification_review_events']:raise ValueError('분류확인 건수가 일치하지 않습니다.')
            frame.attrs['collection_classification_csv_gzip']=gzip.compress(review.to_csv(index=False,lineterminator='\n').encode('utf-8-sig'))
    return frame

def manifest(frame=None, quality_path="data/history_quality.json"):
    if frame is not None:
        return frame.attrs.get(ATTR,{})
    path=Path(quality_path)
    if path.exists():
        return json.loads(path.read_text(encoding="utf-8")).get(ATTR,{})
    return {}

def description(m):
    if not m:return "수집 원장 버전 미확인 · 관리자 업로드 출처를 확인하세요."
    return (f"{m['source']} | {m['version']} | 원천 {m['source_rows']:,}행 · "
            f"운영 {m['operational_rows']:,}행 · 학습 {m['training_rows']:,}행 | "
            f"학습 최종일 {m['data_end_date']} | 완료 {m['completed_tasks']}/{m['total_tasks']}")

def add_excel_metadata(buffer, m):
    if not m:return buffer
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment,Font
    book=load_workbook(buffer)
    sheet=book.create_sheet("데이터출처")
    entries=[("데이터 출처·버전",description(m)),("품질검토 기준일",m['as_of']),
             ("지역 원천값",m['region_basis']),("미해결 작업",m['unresolved']),
             ("예가·1순위 숫자(원천행)",m['numeric_status']['both_numeric']),
             ("예가만 숫자(원천행)",m['numeric_status']['target_only_numeric']),
             ("예가 숫자 없음(원천행)",m['numeric_status']['target_not_numeric']),
             ("품질경고(중복 가능)",f"공고명 누락 {m['warnings']['missing_name']} / 기관 누락 {m['warnings']['missing_org']} / 날짜 미확인 {m['warnings']['missing_date']} / 미래 개찰 {m['warnings']['future_opening']} / 비정상 예가 {m['warnings']['abnormal_target']} / 비정상 1순위 {m['warnings']['abnormal_winner']}"),
             ("중복관리",f"용역/전기공사 {m['dual_profile_events']}키는 사건당 1회 학습; 원천 출처 별도 보존")]
    if m.get('classification_review_events'):
        entries.append(('업종·분류 확인대상',f"두 원문 업종의 기존 분류결과가 다른 {m['classification_review_events']}키(학습 후보 {m['classification_review_learning_rows']}건); 용역 출처 대표값 정책 적용, 실제 업종 원문확인 필요"))
    for a,b in entries:sheet.append([a,b])
    sheet.column_dimensions['A'].width=30;sheet.column_dimensions['B'].width=100
    for index,row in enumerate(sheet,1):
        for cell in row:cell.alignment=Alignment(wrap_text=True,vertical="top")
        row[0].font=Font(bold=True)
        sheet.row_dimensions[index].height=42 if len(str(row[1].value))>80 else 26
    out=io.BytesIO();book.save(out);out.seek(0);return out
