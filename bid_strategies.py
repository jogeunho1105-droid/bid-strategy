"""Versioned observational strategies. No external requests or production writes.
Historical hit rates are descriptive, NOT a current notice's win probability.
"""
from __future__ import annotations
from collections import defaultdict
from datetime import datetime, timezone
import hashlib
import json
import math
import re
import numpy as np
import pandas as pd

VERSION = "v2.15.11"
EFFECTIVE_FROM = "2026-09-28"
AS_OF = "2026-09-27"
LABELS = {"S1":"현행 기준","S2":"기관·분야 축소중심","S3":"시간가중 분위수","S4":"가격구간 중심대체","S5":"가격구간 순차선정"}
MODES = {"분야별 관측 1위 · 시험운영":"auto", **{f"{s} · {v}":s for s,v in LABELS.items()}}
EVIDENCE = {
    "general":{"strategy":"S3","wins":228,"n":4304,"source":"기존사업 주평가","confirmed":False},
    "electric_construction":{"strategy":"S4","wins":573,"n":6538,"source":"전기공사 시장참고","confirmed":False},
    "construction_management":{"strategy":"S3","wins":489,"n":943,"source":"건설사업관리 시장참고","confirmed":False},
}
GRID = np.round(np.arange(-3.,3.0001,.02),4)


def notice_parts(value):
    text=str(value or "").strip().upper()
    match=re.fullmatch(r"([A-Z0-9][A-Z0-9_-]*)-(\d+)",text)
    return (match[1],int(match[2])) if match else (text,None)


def latest_notices(bids):
    """CURRENT inbox only. Never prune historical rows using future amendments.
    Compare suffixes numerically; keep separate organizations/explicit lots.
    Conflicting highest revisions are held rather than arbitrarily selected.
    """
    items=[dict(b) for b in bids];groups=defaultdict(list)
    for i,b in enumerate(items):
        root,rev=notice_parts(b.get("bid_no"))
        key=(root,re.sub(r"\s+","",str(b.get("org",""))),str(b.get("lot_no",""))) if root else ("missing",i)
        groups[key].append((i,rev))
    removed=[];dropped=set()
    for entries in groups.values():
        revisions=[rev for _,rev in entries if rev is not None]
        if not revisions:continue
        maximum=max(revisions);winners=[i for i,rev in entries if rev==maximum]
        chosen=items[winners[0]].get("bid_no","")
        for i,rev in entries:
            if rev is None or rev<maximum:
                dropped.add(i)
                removed.append({"원본행":i+1,"제외공고번호":items[i].get("bid_no",""),"공고명":items[i].get("name",""),"적용공고번호":chosen,"사유":"정정공고: 숫자 차수가 가장 높은 공고만 적용"})
        def payload(b):return {k:v for k,v in b.items() if k not in {"no","_upload_row"}}
        signatures={json.dumps(payload(items[i]),ensure_ascii=False,sort_keys=True,default=str) for i in winners}
        if len(signatures)>1:
            for i in winners:items[i]["revision_warning"]="동일 최고차수 내용 충돌: 원본 확인 전 산정보류"
        elif len(winners)>1:
            for i in winners[1:]:
                dropped.add(i)
                removed.append({"원본행":i+1,"제외공고번호":items[i].get("bid_no",""),"공고명":items[i].get("name",""),"적용공고번호":chosen,"사유":"동일 최고차수 완전 동일 입력 중복"})
    return [b for i,b in enumerate(items) if i not in dropped],sorted(removed,key=lambda r:r["원본행"])


def resolved_strategy(family,mode,when):
    if mode not in ("auto",*LABELS):raise ValueError("지원하지 않는 전략입니다.")
    if mode!="auto":return mode
    if pd.Timestamp(when).normalize()<pd.Timestamp(EFFECTIVE_FROM):return "S1"
    return EVIDENCE.get(family,EVIDENCE["general"])["strategy"]


def _finite(value):
    try:return math.isfinite(float(value))
    except (TypeError,ValueError):return False


def _window(records,date,days,cap):
    lower=date-pd.Timedelta(days=days)
    return [r for r in records if lower<=r["date"]<date][-cap:]


def _winner_records(records):
    return [r for r in records if _finite(r.get("winner")) and abs(float(r["winner"]))<10]


def _weights(records,date):
    return np.array([2.**(-(date-r["date"]).days/180.) for r in records])


def quantiles(family,issuer,date,count):
    """Frozen S3 empirical CDF; 1/2/3 companies use Q50/Q30,70/Q20,50,80."""
    mix=len(issuer)/(len(issuer)+60.);arrays=[];weights=[]
    for records,fraction in ((family,1-mix),(issuer,mix)):
        if not records or fraction<=0:continue
        w=_weights(records,date);arrays.extend(float(r["value"]) for r in records);weights.extend(w/w.sum()*fraction)
    if not arrays:return []
    values=np.asarray(arrays);w=np.asarray(weights);order=np.argsort(values);values=values[order];w=w[order]
    cumulative=np.cumsum(w);q={1:[.5],2:[.3,.7],3:[.2,.5,.8]}[count]
    return [round(float(values[min(np.searchsorted(cumulative,x*cumulative[-1]),len(values)-1)]),4) for x in q]


def robust_center(records):
    a=np.asarray([float(r["value"]) for r in records])
    if not len(a):return None
    lo,hi=np.quantile(a,[.1,.9]);trimmed=a[(a>=lo)&(a<=hi)]
    return float(.6*np.median(a)+.4*np.mean(trimmed))


def shrink_center(family,issuer,date,fallback):
    f=_window(family,date,365,200);local=_window(issuer,date,365,50)
    if len(f)<20:return fallback
    prior=robust_center(f);mix=len(local)/(len(local)+20.)
    value=mix*(robust_center(local) if local else prior)+(1-mix)*prior
    lo,hi=np.quantile([r["value"] for r in f],[.05,.95])
    return round(float(np.clip(value,lo,hi)),4)


def interval_point(family,issuer,date,retained,fallback):
    def score(records):
        valid=_winner_records(records)
        if not valid:return np.zeros(len(GRID)),0
        lower=np.array([r["value"] for r in valid]);upper=np.array([r["winner"] for r in valid]);weights=_weights(valid,date)
        hits=(lower[:,None]<GRID)&(GRID<upper[:,None])
        cumulative=np.cumsum(np.pad(hits.astype(float),((0,0),(6,5))),axis=1)
        smooth=(cumulative[:,11:]-cumulative[:,:-11])/11.
        covered=np.zeros(len(lower),dtype=bool)
        for rate in retained:covered|=(lower<rate)&(rate<upper)
        return np.einsum("i,ij->j",weights*(~covered),smooth)/weights.sum(),len(valid)
    global_score,n=score(family);local_score,m=score(issuer)
    if len(family)<50 or n<30:return fallback
    mix=m/(m+60.);scores=mix*local_score+(1-mix)*global_score
    if not np.any(scores>0):return fallback
    best=np.flatnonzero(np.isclose(scores,scores.max(),atol=1e-12))
    return float(GRID[best[np.argmin(abs(GRID[best]-fallback))]])


def calculate_selected(family_all,issuer_all,date,count,selected,baseline):
    """730d -> same-scope 1460d -> S3 -> reference >=5 -> explicit hold.
    Caller guarantees compatible model family and KEPCO scope.
    """
    notes=[];used=selected;days=730
    family=_window(family_all,date,days,2000);issuer=_window(issuer_all,date,days,300)
    def enough(f):
        return (len(f)>=50 and len(_winner_records(f))>=30) if selected in ("S4","S5") else len(f)>=30
    if not enough(family):
        extended=_window(family_all,date,1460,2000)
        if len(extended)>len(family):
            family=extended;issuer=_window(issuer_all,date,1460,300);days=1460
            notes.append("최근2년 표본 부족 → 동일 모델군·입찰범위 최대4년 확장")
    meta={"requested_strategy":selected,"selected_strategy":selected,"window_days":days,"family_n":len(family),"issuer_n":len(issuer),"winner_n":len(_winner_records(family)),"recent90_n":len(_window(issuer_all,date,90,1000000)),"family_recent90_n":len(_window(family_all,date,90,1000000)),"last_same_issuer_date":str(issuer[-1]["date"].date()) if issuer else None,"sample_status":"기본","fallback_notes":notes}
    w=_weights(family,date);meta["effective_n"]=round(float(w.sum()**2/(w@w)),2) if len(w) else 0.
    if not meta["recent90_n"]:notes.append("동일기관·모델군 최근90일 0건: 최신 기관추세로 해석 금지")
    if len(family)<5:
        meta.update(selected_strategy="HOLD",sample_status="산정보류");notes.append("동일 범위 유효표본 5건 미만: 임의 추천값 생성 안 함")
        return [],meta
    if len(family)<30 or meta["effective_n"]<5:
        value=round(robust_center(issuer if len(issuer)>=5 else family),4)
        meta.update(selected_strategy="REFERENCE",sample_status="자료부족 참고값")
        notes.append("소표본 강건중심 참고값; 복수업체 값이 같아 분산효과 없음; 성능 미검증")
        return [value]*count,meta
    if selected in ("S4","S5") and (not enough(family) or not any(r["winner"]>r["value"] for r in _winner_records(family))):
        used="S3";notes.append("1순위 가격구간 표본 부족 → 예가분포 S3 대체; 원전략 통과율 승계 안 함")
    if selected=="S2" and len(_window(family,date,365,200))<20:
        used="S3";notes.append("S2 최근1년 표본 부족 → S3 대체; 원전략 통과율 승계 안 함")
    ci=0 if count==1 else 1;base=list(baseline or quantiles(family,issuer,date,count))
    if len(base)!=count:base=quantiles(family,issuer,date,count)
    if used=="S3":values=quantiles(family,issuer,date,count)
    elif used=="S2":
        values=list(base);values[ci]=shrink_center(family,issuer,date,base[ci])
    elif used=="S4":
        values=list(base);values[ci]=interval_point(family,issuer,date,[v for i,v in enumerate(base) if i!=ci],base[ci])
    elif used=="S5":
        values=[]
        for i in range(count):
            fallback=base[ci] if i==0 else base[0] if i==1 else base[2]
            values.append(interval_point(family,issuer,date,values,fallback))
    else:values=base
    if len(set(values))<count:notes.append("업체 간 동일 추천값 존재: 독립 분산효과로 해석하지 않음")
    meta["selected_strategy"]=used
    meta["sample_status"]="대체·확인필요" if days!=730 or used!=selected else "관측1위·시험" if selected in ("S3","S4") else "연구전략·시험"
    return [round(float(v),4) for v in values],meta


def _blocked(bid,row,scope):
    if bid.get("revision_warning"):return bid["revision_warning"]
    if re.search(r"\[취소\]|취소공고|공고취소",str(bid.get("name",""))) or bid.get("cancelled") is True:
        return "최고차수 공고가 취소 상태: 하위 차수를 되살리지 않음"
    if row["family"]=="other":return "미분류 공고: 관련 없는 업종의 값으로 대체하지 않음"
    if row["family"]=="kepco_supervision" and "단가" in row["name"]:
        return "한전 단가감리: 단가를 총 추정가격으로 간주할 수 없어 원문 범위·총액 확인 전 산정보류"
    if row["family"]=="kepco_supervision" and row["base"]<=0:
        return "한전 감리 금액 미확인: 지역/전국·업체 수 확인 전 산정보류"
    return ""


def recommend_bids(bids,history,policy=None,mode="auto"):
    """One chronological shared engine; no same-day/future training."""
    import bid_engine as engine
    import bid_rules as rules
    policy=engine.load_policy() if policy is None else policy
    if history is None or len(history)==0:
        return [dict(rates=[],status="산정보류",error="유효 낙찰이력 없음",scope=rules.classify_bid_scope(b)) for b in bids]
    prepared=engine.prepare_history(history);records=[engine.source_record(s) for s in prepared.to_dict("records")]
    history_hash=hashlib.sha256(prepared[["_date","공고번호","_value","_winner"]].to_csv(index=False).encode()).hexdigest()
    targets=[];output=[None]*len(bids)
    for i,b in enumerate(bids):
        try:targets.append((i,engine.bid_record(b)))
        except ValueError:output[i]=dict(rates=[],status="산정보류",error="투찰마감일 미확인",scope=rules.classify_bid_scope(b))
    instance=engine.PredictionEngine();cursor=0
    for i,row in sorted(targets,key=lambda v:v[1]["date"]):
        while cursor<len(records) and records[cursor]["date"]<row["date"]:
            instance.add(records[cursor]);cursor+=1
        scope=engine.scope_for(row);reason=_blocked(bids[i],row,scope)
        if reason:
            output[i]=dict(rates=[],status="산정보류",error=reason,scope=scope,selected_strategy="HOLD",history_hash=history_hash);continue
        selected=resolved_strategy(row["family"],mode,row["date"])
        baseline=instance.predict(row,alternatives=True)
        active=not policy.get("effective_from") or row["date"]>=pd.Timestamp(policy["effective_from"])
        key=row["family"]+"|"+scope["scope"];route=policy.get("rules",{}) if active else {}
        old_strategy=route.get(key+"|"+row["org"],route.get(key,"current"))
        base_rates=baseline["recommendations"].get(old_strategy,baseline["recommendations"]["current"]) if baseline else []
        poolkey=engine.pool_key(row);family=list(instance.family[poolkey]);issuer=list(instance.issuer_family[(poolkey,row["org"])])
        root,_=notice_parts(bids[i].get("bid_no"))
        def same_event(r):return bool(root) and r["org"]==row["org"] and notice_parts(r.get("notice"))[0]==root
        contaminated=any(same_event(r) for r in family)
        if contaminated and selected in ("S1","S2","S4","S5"):
            output[i]=dict(rates=[],status="산정보류",error="동일 공고의 과거 정정차수 결과 존재: 자기결과 학습 방지를 위해 확인 필요",scope=scope,selected_strategy="HOLD",history_hash=history_hash);continue
        family=[r for r in family if not same_event(r)];issuer=[r for r in issuer if not same_event(r)]
        if selected=="S1":
            values=list(base_rates)
            meta={"requested_strategy":"S1","selected_strategy":"S1","sample_status":"현행 기준","family_n":len(_window(family,row["date"],730,2000)),"issuer_n":len(_window(issuer,row["date"],730,300)),"recent90_n":len(_window(issuer,row["date"],90,1000000)),"window_days":730,"fallback_notes":[],"winner_n":len(_winner_records(_window(family,row["date"],730,2000)))}
            if not values:
                values,meta=calculate_selected(family,issuer,row["date"],scope["company_count"],"S3",[])
                meta["fallback_notes"].insert(0,"S1 산출 불가 → 동일 범위 분포 대체")
        else:values,meta=calculate_selected(family,issuer,row["date"],scope["company_count"],selected,base_rates)
        if scope.get("scope_assumed"):meta["fallback_notes"].append("입찰범위·참가회사 수에 기존 운영 가정 포함: 참가자격 별도 확인")
        if row["family"]=="kepco_supervision" and any("단가" in r["name"] for r in family):
            meta["fallback_notes"].append("학습범위에 단가감리 이력 포함: 기존 분류 정합성 검토 필요")
        subtype=[r for r in issuer if r["svc"]==row["svc"]];meta["issuer_subtype90_n"]=len(_window(subtype,row["date"],90,1000000))
        if meta["issuer_subtype90_n"]<5:meta["fallback_notes"].append(f"동일기관·세부업종 최근90일 {meta['issuer_subtype90_n']}건: 충분한 최신 추세 근거 없음")
        result=dict(baseline or {},rates=values,scope=scope,strategy=old_strategy if selected=="S1" else meta["selected_strategy"],baseline_rates=base_rates,status=meta["sample_status"] if values else "산정보류",error="" if values else "동일 범위 이력 부족: 산정보류",history_hash=history_hash,latest_train_date=str(instance.last_date.date()) if instance.last_date is not None else None,information_cutoff=str(row["date"].date()),cutoff_assumption="투찰마감일 전일까지만 사용",policy_as_of=AS_OF,policy_effective_from=EFFECTIVE_FROM,selection_mode=mode,evidence=EVIDENCE.get(row["family"],EVIDENCE["general"]),**meta)
        output[i]=result
    return output


def recommendation_records(bid,prediction):
    if not prediction or not prediction.get("rates"):return []
    selected=prediction.get("selected_strategy","S1");label=LABELS.get(selected,"소표본 참고값");rates=prediction["rates"]
    if len(rates)!=prediction["scope"]["company_count"] or any(not _finite(x) for x in rates):raise ValueError("추천 수·유한값 검사 실패")
    notes="; ".join(prediction.get("fallback_notes",[]))
    basis=(f"{label}; 동일모델군 {prediction.get('family_n',0)}건 / 동일기관 {prediction.get('issuer_n',0)}건; "
           f"최근90일 n={prediction.get('recent90_n',0)}; 과거 {prediction.get('window_days',730)}일; "
           f"{prediction.get('sample_status','확인필요')}; {notes}; 관측 통과율은 해당 공고의 낙찰확률 아님")
    return [dict(company=f"업체 {i+1}",rate=float(v),role=label,model=selected,model_label=f"{selected} {label}",basis_n=int(prediction.get("family_n",0)),basis=basis,strategy=selected,alternative=selected!="S1") for i,v in enumerate(rates)]


def snapshot(results,removed=()):
    data={"model_version":VERSION,"created_at_utc":datetime.now(timezone.utc).isoformat(),"policy_as_of":AS_OF,"effective_from":EFFECTIVE_FROM,"removed_revisions":list(removed),"results":[{"bid":r.get("bid"),"prediction":r.get("prediction"),"recommendations":r.get("recommendations"),"error":r.get("error","")} for r in results],"warning":"관측 순위 기반 시험운영; 실제 낙찰확률 아님; 실제 마감시각·자격·금액산식 별도 확인"}
    return json.dumps(data,ensure_ascii=False,indent=2,default=str,allow_nan=False).encode("utf-8")


def add_review_sheets(buffer,results,removed=()):
    """Preserve existing sheets; append strategy/revision audit sheets."""
    import io
    from openpyxl import load_workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    wb=load_workbook(io.BytesIO(buffer.getvalue()));ws=wb.create_sheet("전략선택근거")
    headers=["공고번호","공고명","요청전략","실제적용전략","상태","모델군표본","기관표본","최근90일기관표본","최근90일기관세부업종","사용기간일","유효표본수","학습최종일","정보컷오프","대체·경고","모델버전"]
    ws.append(headers)
    for row in results:
        p=row.get("prediction") or {};b=row.get("bid") or {}
        ws.append([b.get("bid_no"),b.get("name"),p.get("requested_strategy"),p.get("selected_strategy"),p.get("status"),p.get("family_n"),p.get("issuer_n"),p.get("recent90_n"),p.get("issuer_subtype90_n"),p.get("window_days"),p.get("effective_n"),p.get("latest_train_date"),p.get("information_cutoff"),"; ".join([p.get("error","")]+p.get("fallback_notes",[])),VERSION])
    audit=wb.create_sheet("정정공고제외");keys=["원본행","제외공고번호","공고명","적용공고번호","사유"];audit.append(keys)
    for row in removed:audit.append([row.get(k) for k in keys])
    for sheet in (ws,audit):
        sheet.freeze_panes="C2";sheet.auto_filter.ref=sheet.dimensions;sheet.sheet_view.showGridLines=False
        for cell in sheet[1]:
            cell.font=Font(name="맑은 고딕",bold=True,color="FFFFFF");cell.fill=PatternFill("solid",fgColor="1A2744")
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width=22
            for cell in column:
                if isinstance(cell.value,str) and cell.value.startswith("="):cell.data_type="s"
                cell.alignment=Alignment(vertical="top",wrap_text=True)
        sheet.column_dimensions["B"].width=42
        for row in range(2,sheet.max_row+1):sheet.row_dimensions[row].height=42
    ws.column_dimensions["N"].width=75
    out=io.BytesIO();wb.save(out);out.seek(0);return out
